"""
Parse updated earring list.xlsx → earrings-import.json
CRITICAL: keep name + description EXACTLY as in the sheet (no rewriting).
"""
from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

SRC = Path(r"C:\Users\rohan\Downloads\updated earring list.xlsx")
OUT = Path(r"c:\Projects\Velvel_shell\docs\earrings-import.json")

# Letter → metal when product copy mentions both gold plating + 925 silver
FINISH_A_GOLD = "18K Gold-Plated"
FINISH_SILVER = "925 Silver"

# Letter → gem/stone colour when product has multi-stone colourways
LETTER_STONE = {
    "A": ("Option A", "#C5C5C5"),
    "B": ("Option B", "#C5C5C5"),
    "C": ("Option C", "#C5C5C5"),
    "D": ("Option D", "#C5C5C5"),
    "E": ("Option E", "#C5C5C5"),
}

# Stronger stone maps when title/desc names the stones
STONE_SERIES = [
    (
        r"aquamarine",
        [("Blue", "#4A7BA7"), ("Green", "#2F6B4F"), ("Clear", "#E8EEF2")],
    ),
    (
        r"pink\s*white\s*crystal|flower in the rain",
        [
            ("Pink", "#E8B4C4"),
            ("White", "#F5F5F5"),
            ("Clear", "#E8EEF2"),
            ("Rose", "#D4A5A5"),
        ],
    ),
    (
        r"garnet|topaz|amethyst|peridot",
        [
            ("Garnet", "#A83B3B"),
            ("Topaz", "#D4AF64"),
            ("Amethyst", "#6B3FA0"),
            ("Peridot", "#7A9E7E"),
        ],
    ),
    (
        r"\bagate\b",
        [("White", "#F5F5F5"), ("Pink", "#E8B4C4"), ("Green", "#2F6B4F")],
    ),
    (
        r"pearl",
        [("White", "#F5F5F5"), ("Pink", "#E8B4C4"), ("Purple", "#6B3FA0")],
    ),
    (
        r"amber",
        [("Amber", "#D4AF64"), ("Honey", "#C9A227")],
    ),
    (
        r"crystal",
        [("Clear", "#E8EEF2"), ("Pink", "#E8B4C4"), ("White", "#F5F5F5")],
    ),
]


def slugify(text: str) -> str:
    s = re.sub(r"[^a-z0-9]+", "-", (text or "").lower()).strip("-")
    return s[:80] or "item"


def parse_name_desc_exact(cell) -> tuple[str, str]:
    """Split first line = name, rest = description. Do NOT rewrite text."""
    if not cell:
        return "", ""
    raw = str(cell).replace("\r\n", "\n").replace("\r", "\n").strip()
    if not raw:
        return "", ""
    lines = raw.split("\n")
    name = lines[0].strip()
    desc = "\n".join(lines[1:]).strip() if len(lines) > 1 else ""
    return name, desc


def finish_flags(text: str) -> tuple[bool, bool]:
    t = (text or "").lower()
    gold = bool(re.search(r"18\s*k|gold[\s-]*plat|gold plating", t))
    silver = bool(re.search(r"925|sterling", t))
    return gold, silver


def parse_price(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        n = int(v)
        return n if 500 <= n <= 50000 else None
    s = str(v).strip().replace(",", "")
    if not s or s.startswith("#"):
        return None
    try:
        n = int(float(s))
    except ValueError:
        return None
    return n if 500 <= n <= 50000 else None


def base_key(sku: str) -> tuple[str, str]:
    sku = sku.strip().upper()
    if sku.startswith("YWJE"):
        return sku, "DEFAULT"
    m = re.match(r"^(LFJ[A-Z]\d+)([A-Z])(\d+)$", sku)
    if m:
        return m.group(1), f"{m.group(2)}{m.group(3)}"
    m = re.match(r"^(LFJ[A-Z]\d+)([A-Z])$", sku)
    if m:
        return m.group(1), m.group(2)
    m = re.match(r"^(LFJ[A-Z]\d+)$", sku)
    if m:
        return m.group(1), "DEFAULT"
    m = re.match(r"^(.+?)([A-Z])$", sku)
    if m:
        return m.group(1), m.group(2)
    return sku, "DEFAULT"


def assign_finish(
    token: str,
    gold: bool,
    silver: bool,
    sibling_letters: set[str],
    family_tokens: list[str],
) -> str:
    letter_m = re.match(r"^([A-Z])", token or "A")
    letter = letter_m.group(1) if letter_m else "A"

    digit_series = [t for t in family_tokens if re.match(r"^[A-Z]\d+$", t)]
    if len(digit_series) >= 2:
        # Colour series on one metal
        return FINISH_A_GOLD if gold else FINISH_SILVER

    if gold and silver and len(sibling_letters) >= 2:
        # Dual metal: A = gold, other letters = silver
        if letter == "A":
            return FINISH_A_GOLD
        return FINISH_SILVER

    if gold and not silver:
        return FINISH_A_GOLD
    if gold and silver and len(sibling_letters) < 2:
        # Single SKU but copy mentions both — prefer gold if A, else as written
        return FINISH_A_GOLD if letter == "A" else FINISH_SILVER
    return FINISH_SILVER


def assign_stone(
    token: str,
    text: str,
    family_tokens: list[str],
    dual_finish: bool,
) -> tuple[str, str, bool]:
    """
    Return (colorLabel, swatch, isGemAxis).
    When dual finish only (A/B gold/silver), colour stays Default.
    When multi stone SKUs, map to gem colours.
    """
    digit_tokens = [t for t in family_tokens if re.match(r"^[A-Z]\d+$", t)]
    m = re.match(r"^([A-Z])(\d+)$", token or "")

    if m and len(digit_tokens) > 1:
        letters = sorted({t[0] for t in digit_tokens})
        if len(letters) == 1:
            idx = int(m.group(2)) - 1
            label_key = m.group(2)
        else:
            idx = letters.index(m.group(1)) if m.group(1) in letters else 0
            label_key = m.group(1)
        for pat, colors in STONE_SERIES:
            if re.search(pat, text, re.I):
                if 0 <= idx < len(colors):
                    return colors[idx][0], colors[idx][1], True
        return f"Colour {label_key}", "#C5C5C5", True

    letters = sorted(
        {
            re.match(r"^([A-Z])", t).group(1)
            for t in family_tokens
            if re.match(r"^[A-Z]", t)
        }
    )

    # Dual finish A/B (or A/E) without digit series → metal axis, not stone
    if dual_finish and len(family_tokens) <= 2 and not digit_tokens:
        return "Default", "#C5C5C5", False

    # 3+ letter variants → treat as stone/colour options
    if len(letters) >= 3:
        letter = token[0] if token and token[0].isalpha() else "A"
        li = letters.index(letter) if letter in letters else 0
        for pat, colors in STONE_SERIES:
            if re.search(pat, text, re.I):
                if li < len(colors):
                    return colors[li][0], colors[li][1], True
        label, swatch = LETTER_STONE.get(letter, (f"Option {letter}", "#C5C5C5"))
        return label, swatch, True

    # Two letter variants, NOT dual metal → colour options (Option A / Option B)
    if len(letters) == 2 and not dual_finish:
        letter = token[0] if token and token[0].isalpha() else "A"
        for pat, colors in STONE_SERIES:
            if re.search(pat, text, re.I):
                li = letters.index(letter) if letter in letters else 0
                if li < len(colors):
                    return colors[li][0], colors[li][1], True
        label, swatch = LETTER_STONE.get(letter, (f"Option {letter}", "#C5C5C5"))
        return label, swatch, True

    return "Default", "#C5C5C5", False


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb["Sheet1"]

    raw_rows = []
    pending_name = ""
    pending_desc = ""
    pending_raw = ""
    pending_drive = ""

    for r in range(1, ws.max_row + 1):
        sku_cell = ws.cell(r, 1).value
        if not sku_cell:
            continue
        sku = str(sku_cell).strip().upper()
        cell3 = ws.cell(r, 3).value
        name, desc = parse_name_desc_exact(cell3)
        if name:
            pending_name, pending_desc = name, desc
            pending_raw = str(cell3)
        drive = ws.cell(r, 7).value
        if drive:
            pending_drive = str(drive).strip()
        price = parse_price(ws.cell(r, 6).value)
        stock_hint = ws.cell(r, 5).value
        stock = 10
        if isinstance(stock_hint, (int, float)) and 0 < stock_hint < 100:
            stock = int(stock_hint)

        if sku.startswith("YWJE"):
            continue

        raw_rows.append(
            {
                "sku": sku,
                "name": pending_name,
                "desc": pending_desc,
                "raw": pending_raw,
                "price": price,
                "stock": stock,
                "drive": pending_drive,
            }
        )

    families = defaultdict(list)
    for row in raw_rows:
        fam, token = base_key(row["sku"])
        families[fam].append({**row, "token": token})

    products = []
    skipped = []

    for fam, items in families.items():
        name = next((i["name"] for i in items if i["name"]), "")
        desc = next((i["desc"] for i in items if i["desc"]), "")
        raw = next((i["raw"] for i in items if i.get("raw")), "")
        if not name:
            skipped.append({"fam": fam, "reason": "no name"})
            continue

        priced = [i for i in items if i["price"] is not None]
        if not priced:
            skipped.append(
                {
                    "fam": fam,
                    "reason": "no valid price",
                    "skus": [i["sku"] for i in items],
                }
            )
            continue

        text = f"{name}\n{desc}\n{raw}"
        gold, silver = finish_flags(text)
        tokens = [i["token"] for i in priced]
        letters = {
            re.match(r"^([A-Z])", t).group(1)
            for t in tokens
            if re.match(r"^[A-Z]", t)
        }
        digit_series = [t for t in tokens if re.match(r"^[A-Z]\d+$", t)]
        dual_finish = bool(gold and silver and len(letters) >= 2 and len(digit_series) < 2)

        variants = []
        color_used = False
        finishes_used = set()

        for i in priced:
            finish = assign_finish(i["token"], gold, silver, letters, tokens)
            color, swatch, is_gem = assign_stone(i["token"], text, tokens, dual_finish)
            if is_gem:
                color_used = True
            finishes_used.add(finish)

            # Exact name/desc on every variant (for PDP copy)
            variants.append(
                {
                    "sku": i["sku"],
                    "token": i["token"],
                    "colorLabel": color,
                    "colorSwatch": swatch,
                    "finishLabel": finish,
                    "material": finish,
                    "price": i["price"],
                    "stock": i["stock"],
                    "description": desc if desc else "",
                    "heroUrl": "",
                    "gallery": [],
                    "isPrimary": False,
                    "features": [f"SKU {i['sku']}", finish],
                    "sizeTags": ["One size"],
                    "length": "One size",
                    "reviews": {"rating": 0, "count": 0, "highlights": []},
                    "active": True,
                }
            )

        # For dual-finish with Default colour, expose Gold/Silver as selectable
        # via color_label workaround if finish_label column missing — still store finishLabel
        if dual_finish and not color_used:
            for v in variants:
                if v["finishLabel"] == FINISH_A_GOLD:
                    v["colorLabel"] = "Gold"
                    v["colorSwatch"] = "#D4AF64"
                else:
                    v["colorLabel"] = "Silver"
                    v["colorSwatch"] = "#C5C5C5"
            # Storefront treats Gold/Silver as metal not gem; hasColorVariants helps listing
            color_used = False

        for v in variants:
            v["isPrimary"] = False
        primary_set = False
        for v in variants:
            if v["finishLabel"] == FINISH_A_GOLD:
                v["isPrimary"] = True
                primary_set = True
                break
        if not primary_set and variants:
            variants[0]["isPrimary"] = True

        has_finish = len({v["finishLabel"] for v in variants}) > 1
        has_color = len({v["colorLabel"] for v in variants if v["colorLabel"] not in ("Default", "Gold", "Silver")}) > 1
        # Also colour variants if Option A/B or named stones
        if not has_color:
            labels = {v["colorLabel"] for v in variants}
            if len(labels) > 1 and not labels <= {"Gold", "Silver", "Default"}:
                has_color = True
            elif any(
                v["colorLabel"] not in ("Default", "Gold", "Silver") for v in variants
            ) and len(variants) > 1:
                has_color = True

        # Dual finish only → hasFinishVariants; show Gold/Silver as colour tabs in DB workaround
        if has_finish and not has_color:
            # Keep Gold/Silver labels for selector when finish column absent
            pass

        products.append(
            {
                # EXACT sheet name — never cleaned
                "name": name,
                "slug": slugify(f"{name}-{fam}"),
                "category": "earrings",
                "collection": "the-shell-edit",
                "badge": "",
                "status": "active",
                "hasColorVariants": has_color or (has_finish and any(v["colorLabel"] in ("Gold", "Silver") for v in variants)),
                "hasFinishVariants": has_finish,
                "familyKey": fam,
                "drive": next((i["drive"] for i in items if i.get("drive")), ""),
                "variants": variants,
            }
        )

    # Unique slugs
    seen = {}
    for p in products:
        base = p["slug"]
        if base in seen:
            seen[base] += 1
            p["slug"] = f"{base}-{seen[base]}"
        else:
            seen[base] = 1

    OUT.write_text(
        json.dumps({"products": products, "skipped": skipped}, indent=2, ensure_ascii=False),
        encoding="utf-8",
    )
    print(f"Wrote {len(products)} products, skipped {len(skipped)}")
    fc = Counter()
    cc = Counter()
    for p in products:
        for v in p["variants"]:
            fc[v["finishLabel"]] += 1
            cc[v["colorLabel"]] += 1
    print("finishes", dict(fc))
    print("colours", dict(cc))
    # Prove name/desc untouched for first product
    if products:
        p0 = products[0]
        print("SAMPLE NAME:", repr(p0["name"][:80]))
        print("SAMPLE DESC:", repr((p0["variants"][0]["description"] or "")[:120]))


if __name__ == "__main__":
    main()
