"""Parse updated-earring-list.xlsx → clean product families JSON for import."""
from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

SRC = Path(r'c:\Projects\Velvel_shell\docs\updated-earring-list.xlsx')
OUT = Path(r'c:\Projects\Velvel_shell\docs\earrings-import.json')

# LotusFun-style trailing letter → metal when both gold + silver appear in copy
# and the pair is A vs B/E/D (A usually gold / higher price).
LETTER_FINISH = {
    'A': '18K Gold-Plated',
    'B': '925 Silver',
    'C': '925 Silver',
    'D': '925 Silver',
    'E': '925 Silver',
}

GEM_FROM_LETTER = {
    'A': ('Default', '#C5C5C5'),
    'B': ('Default', '#C5C5C5'),
    'C': ('Default', '#C5C5C5'),
    'D': ('Default', '#C5C5C5'),
    'E': ('Default', '#C5C5C5'),
}

# Known multi-stone series from titles
STONE_SERIES = [
    (r'aquamarine', [('Blue', '#4A7BA7'), ('Green', '#2F6B4F'), ('Clear', '#E8EEF2')]),
    (r'pink\s*white\s*crystal|flower in the rain', [
        ('Pink', '#E8B4C4'),
        ('White', '#F5F5F5'),
        ('Clear', '#E8EEF2'),
        ('Rose', '#D4A5A5'),
    ]),
    (r'agate', [('White', '#F5F5F5'), ('Pink', '#E8B4C4'), ('Green', '#2F6B4F')]),
    (r'garnet|topaz|amethyst|peridot', [
        ('Garnet', '#A83B3B'),
        ('Topaz', '#D4AF64'),
        ('Amethyst', '#6B3FA0'),
        ('Peridot', '#7A9E7E'),
    ]),
    (r'pearl', [('White', '#F5F5F5'), ('Pink', '#E8B4C4'), ('Purple', '#6B3FA0')]),
]


def slugify(text: str) -> str:
    s = re.sub(r'[^a-z0-9]+', '-', (text or '').lower()).strip('-')
    return s[:80] or 'item'


def clean_title(raw: str) -> str:
    if not raw:
        return ''
    line = raw.splitlines()[0].strip()
    line = re.sub(
        r'^(Wholesale|Custom Wholesale|Factory Wholesale|Hot\s*Hand Made|New Arrival|'
        r'New Design|Fashion Jewelry|LotusFun Wholesale|Lotusfun Simple Design|'
        r'Lotusfun|Top Elegant|Handmade Luxury|Hand Made|Handmade|Real 925|'
        r'Pure 925|Original|2025 Fashion New Model|Korea New Design|High Quality|'
        r'ED-JS19|Custom Pure|Fashion New|Hot\s+)\s*',
        '',
        line,
        flags=re.I,
    )
    # Drop long material boilerplate from titles
    line = re.sub(
        r'\b(925 Sterling Silver|925 Silver|18k Gold Plated|18K Gold Plated|'
        r'18k Gold Plating|Gold Plating|Sterling Silver)\b',
        '',
        line,
        flags=re.I,
    )
    line = re.sub(
        r'\b(Wholesale|China|Hot Sale|New Model|for Women|Brincos Grandes|'
        r'Jewelry Gift|Prevent Allergy|Fine Jewelry|Luxury Jewelry Gift|'
        r'Elegant Lady|Fashion Jewelry|Jewellery|Jewelry Set|Novelty Chinese|'
        r'Original Design|Personalized French Style)\b',
        '',
        line,
        flags=re.I,
    )
    line = re.sub(r'\s{2,}', ' ', line).strip(' -|,')
    if len(line) > 64:
        m = re.search(r'([A-Za-z][A-Za-z0-9 \-\']{8,60}?Earrings)', line, re.I)
        if m:
            line = m.group(1).strip()
    if not line:
        line = raw.splitlines()[0].strip()[:64]
    return line


def parse_name_desc(cell) -> tuple[str, str]:
    if not cell:
        return '', ''
    raw = str(cell).strip()
    lines = [ln.strip() for ln in raw.splitlines() if ln.strip()]
    if not lines:
        return '', ''
    name = clean_title(lines[0])
    desc = ' '.join(lines[1:]).strip() if len(lines) > 1 else ''
    if not desc and '  ' in lines[0]:
        # sometimes name + desc on one line separated poorly
        desc = ''
    return name, desc


def finish_flags(text: str) -> tuple[bool, bool]:
    t = (text or '').lower()
    gold = bool(re.search(r'18\s*k|gold[\s-]*plat|gold plating', t))
    silver = bool(re.search(r'925|sterling', t))
    return gold, silver


def parse_price(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        n = int(v)
        return n if 500 <= n <= 50000 else None  # skip stock-qty junk
    s = str(v).strip().replace(',', '')
    if not s or s.startswith('#'):
        return None
    try:
        n = int(float(s))
    except ValueError:
        return None
    return n if 500 <= n <= 50000 else None


def base_key(sku: str) -> tuple[str, str]:
    """Return (family_key, variant_token)."""
    sku = sku.strip().upper()
    # YWJE keep whole as family (often unique)
    if sku.startswith('YWJE'):
        return sku, 'DEFAULT'
    # Patterns like LFJA0147E1 / LFJA0087B2 / LFJB0066A2
    m = re.match(r'^(LFJ[A-Z]\d+)([A-Z])(\d+)$', sku)
    if m:
        return m.group(1), f'{m.group(2)}{m.group(3)}'
    m = re.match(r'^(LFJ[A-Z]\d+)([A-Z])$', sku)
    if m:
        return m.group(1), m.group(2)
    m = re.match(r'^(LFJ[A-Z]\d+)$', sku)
    if m:
        return m.group(1), 'DEFAULT'
    # fallback strip trailing letter
    m = re.match(r'^(.+?)([A-Z])$', sku)
    if m:
        return m.group(1), m.group(2)
    return sku, 'DEFAULT'


def assign_finish(
    token: str,
    gold: bool,
    silver: bool,
    price: int | None,
    sibling_prices: dict,
    family_tokens: list[str],
) -> str:
    letter = re.match(r'^([A-Z])', token or 'A')
    letter = letter.group(1) if letter else 'A'
    # Letter+digit series (A2/B1/C1) are usually stone colours on ONE finish
    digit_series = [t for t in family_tokens if re.match(r'^[A-Z]\d+$', t)]
    if len(digit_series) >= 2:
        return '18K Gold-Plated' if gold else '925 Silver'

    if gold and silver:
        # Dual finish product: A → gold, others → silver (common LotusFun)
        if letter == 'A':
            return '18K Gold-Plated'
        if letter in ('B', 'C', 'D', 'E'):
            if letter == 'E' and 'A' not in sibling_prices and 'B' not in sibling_prices:
                return '18K Gold-Plated'
            return '925 Silver'
        return LETTER_FINISH.get(letter, '925 Silver')
    if gold and not silver:
        return '18K Gold-Plated'
    return '925 Silver'


def assign_gem(token: str, name: str, desc: str, family_tokens: list[str]) -> tuple[str, str, bool]:
    """Return colorLabel, swatch, hasColorVariants contribution."""
    text = f'{name} {desc}'.lower()
    # Letter+digit series (A2/B1/C1 or B1/B2/B3)
    m = re.match(r'^([A-Z])(\d+)$', token)
    digit_tokens = [t for t in family_tokens if re.match(r'^[A-Z]\d+$', t)]
    if m and len(digit_tokens) > 1:
        letters = sorted({t[0] for t in digit_tokens})
        if len(letters) == 1:
            idx = int(m.group(2)) - 1
            label_key = m.group(2)
        else:
            idx = letters.index(m.group(1)) if m.group(1) in letters else 0
            label_key = m.group(1)
        for pat, colors in STONE_SERIES:
            if re.search(pat, text):
                if 0 <= idx < len(colors):
                    return colors[idx][0], colors[idx][1], True
        return f'Colour {label_key}', '#C5C5C5', True

    # Letter-only multiples that look like stone colours in anemone etc.
    letters = sorted({re.match(r'^([A-Z])', t).group(1) for t in family_tokens if re.match(r'^[A-Z]', t)})
    if len(letters) >= 3 and not (set(letters) <= {'A', 'B'}):
        # map C/D/E to gem-ish labels when stone words present
        for pat, colors in STONE_SERIES:
            if re.search(pat, text):
                li = letters.index(token[0]) if token and token[0] in letters else 0
                if li < len(colors):
                    return colors[li][0], colors[li][1], True

    # Single or dual letter without clear gem axis → Default colour
    return 'Default', '#C5C5C5', False


def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    ws = wb['Sheet1']

    raw_rows = []
    pending_name = ''
    pending_desc = ''
    pending_drive = ''
    for r in range(1, ws.max_row + 1):
        sku = ws.cell(r, 1).value
        if not sku:
            continue
        sku = str(sku).strip().upper()
        name, desc = parse_name_desc(ws.cell(r, 3).value)
        if name:
            pending_name, pending_desc = name, desc
        drive = ws.cell(r, 7).value
        if drive:
            pending_drive = str(drive).strip()
        price = parse_price(ws.cell(r, 6).value)
        stock_hint = ws.cell(r, 5).value
        # Col E sometimes holds long YWJE specs — ignore non-numeric
        stock = 10
        if isinstance(stock_hint, (int, float)) and 0 < stock_hint < 100:
            stock = int(stock_hint)

        if sku.startswith('YWJE'):
            # Skip noisy wholesale YWJE block for phase-1 earrings cleanup
            continue

        cell3 = ws.cell(r, 3).value
        raw_rows.append({
            'sku': sku,
            'name': pending_name,
            'desc': pending_desc,
            'price': price,
            'stock': stock,
            'drive': pending_drive,
            '_raw': str(cell3 or ''),
        })

    families = defaultdict(list)
    for row in raw_rows:
        fam, token = base_key(row['sku'])
        families[fam].append({**row, 'token': token})

    products = []
    skipped = []
    for fam, items in families.items():
        # Fill name from first named row
        name = next((i['name'] for i in items if i['name']), '')
        desc = next((i['desc'] for i in items if i['desc']), '')
        if not name:
            skipped.append({'fam': fam, 'reason': 'no name'})
            continue
        priced = [i for i in items if i['price'] is not None]
        if not priced:
            skipped.append({'fam': fam, 'reason': 'no valid price', 'skus': [i['sku'] for i in items]})
            continue

        # Use descriptions + original SKU block text (cleaned titles strip metal words)
        blob = ' '.join(
            filter(
                None,
                [name, desc] + [i.get('desc') or '' for i in items] + [i.get('name') or '' for i in items],
            )
        )
        # Prefer raw sheet text for metal detection
        raw_blob = ' '.join(str(i.get('_raw') or '') for i in items)
        gold, silver = finish_flags(f'{blob} {raw_blob}')
        sibling_prices = {}
        for i in priced:
            letter = i['token'][0] if i['token'] and i['token'][0].isalpha() else 'X'
            sibling_prices[letter] = i['price']

        tokens = [i['token'] for i in priced]
        variants = []
        color_used = False
        finishes_used = set()
        for i in priced:
            finish = assign_finish(
                i['token'], gold, silver, i['price'], sibling_prices, tokens,
            )
            color, swatch, is_gem = assign_gem(i['token'], name, desc, tokens)
            if is_gem:
                color_used = True
            finishes_used.add(finish)
            variants.append({
                'sku': i['sku'],
                'token': i['token'],
                'colorLabel': color,
                'colorSwatch': swatch,
                'finishLabel': finish,
                'material': finish,
                'price': i['price'],
                'stock': i['stock'],
                'description': desc or name,
                'heroUrl': '',
                'gallery': [],
                'isPrimary': False,
                'features': [f'SKU {i["sku"]}', 'Handcrafted', finish],
                'sizeTags': ['One size'],
                'length': 'One size',
                'reviews': {'rating': 0, 'count': 0, 'highlights': []},
                'active': True,
            })

        # If dual finish but same Default colour — keep finish axis only
        if not color_used:
            for v in variants:
                v['colorLabel'] = 'Default'
                v['colorSwatch'] = '#C5C5C5'

        # Dedupe only named gem colours that collide (e.g. C1/C2 both Clear).
        # Keep Default×same-finish pairs — those are distinct SKUs (often unknown colourways).
        deduped = []
        seen_keys = set()
        for v in variants:
            key = (v['colorLabel'], v['finishLabel'])
            if (
                v['colorLabel']
                and v['colorLabel'] != 'Default'
                and key in seen_keys
            ):
                for d in deduped:
                    if (d['colorLabel'], d['finishLabel']) == key:
                        d['stock'] = max(d['stock'], v['stock'])
                        break
                continue
            if v['colorLabel'] != 'Default':
                seen_keys.add(key)
            deduped.append(v)
        variants = deduped

        # Distinguish leftover Default pairs by SKU letter so admin/PDP can tell them apart
        default_variants = [v for v in variants if v['colorLabel'] == 'Default']
        if len(default_variants) > 1 and len({v['finishLabel'] for v in default_variants}) == 1:
            for v in default_variants:
                letter = re.sub(r'[^A-Z]', '', (v.get('token') or v.get('sku') or 'X')[-2:])
                letter = (letter[-1] if letter else 'X')
                v['colorLabel'] = f'Option {letter}'
                v['colorSwatch'] = '#C5C5C5'
            # mark as colour variants so both show
            color_used = True

        # Primary = first gold if any else first
        for v in variants:
            v['isPrimary'] = False
        primary_set = False
        for v in variants:
            if v['finishLabel'] == '18K Gold-Plated':
                v['isPrimary'] = True
                primary_set = True
                break
        if not primary_set and variants:
            variants[0]['isPrimary'] = True

        finishes_used = {v['finishLabel'] for v in variants}
        has_finish = len(finishes_used) > 1
        has_color = color_used and len({v['colorLabel'] for v in variants}) > 1

        products.append({
            'name': name,
            'slug': slugify(name) if len(families) else slugify(fam),
            'category': 'earrings',
            'collection': 'the-shell-edit',
            'badge': '',
            'status': 'active',
            'hasColorVariants': has_color,
            'hasFinishVariants': has_finish,
            'familyKey': fam,
            'drive': next((i['drive'] for i in items if i.get('drive')), ''),
            'variants': variants,
        })

    # Ensure unique slugs
    seen = {}
    for p in products:
        base = p['slug']
        if base in seen:
            seen[base] += 1
            p['slug'] = f'{base}-{seen[base]}'
        else:
            seen[base] = 1
        # Prefer family key in slug for stability
        p['slug'] = slugify(f"{p['name']}-{p['familyKey']}")

    OUT.write_text(json.dumps({'products': products, 'skipped': skipped}, indent=2, ensure_ascii=False), encoding='utf-8')
    print(f'Wrote {len(products)} products, skipped {len(skipped)} to {OUT}')
    from collections import Counter
    fc = Counter()
    cc = Counter()
    for p in products:
        for v in p['variants']:
            fc[v['finishLabel']] += 1
            cc[v['colorLabel']] += 1
    print('finishes', dict(fc))
    print('colours', dict(cc))


if __name__ == '__main__':
    main()
